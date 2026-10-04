"""
Bulk participant import: builds an event-specific Excel template (its
actual form fields + whichever blocs/items/workshops it has enabled),
parses+validates an uploaded copy into a ParticipantImportBatch for
review, and processes the resulting rows through the exact same
registration pipeline as the event owner's "Nouvelle inscription"
manual-entry feature (dashboard.views_event_owner.event_owner_new_registration)
-- an imported participant ends up identical to a manually-entered one:
same badge/access, same pricing/discount math, same searchability.
"""
from django.db import transaction
from django.utils import timezone

from events.form_validation_service import create_participant_for_event, get_or_create_user_for_manual_registration
from .blocs_service import compute_order
from .models_blocs import BlocItem
from .models_import import ParticipantImportBatch, ParticipantImportRow
from .views_blocs import get_public_bloc_context

STATUS_HEADER = 'Bloc: Statut *'
RESTAURATION_HEADER = 'Bloc: Restauration'
SOCIAL_EVENT_HEADER = 'Bloc: Événement social'
WORKSHOPS_HEADER = 'Ateliers (workshops)'


def build_import_columns(event):
    """
    The ordered list of columns this event's template/import actually
    needs: one per configured form field, then one per bloc that's both
    enabled AND has at least one real item to choose from (mirrors
    get_public_bloc_context's own filtering exactly, so the import never
    offers a bloc the public form itself wouldn't).

    Returns {'form_fields': [...], 'bloc_columns': [...], 'bloc_context': dict or None}.
    bloc_columns entries: {'header', 'bloc_key' ('status'/'restauration'/
    'social_event'/'workshops'), 'required', 'select_mode', 'valid_values': [names]}.
    """
    form_config = event.custom_forms.first()
    form_fields = []
    if form_config:
        for field in form_config.fields_config:
            form_fields.append({
                'name': field.get('name'),
                'label': field.get('label') or field.get('name'),
                'type': field.get('type', 'text'),
                'required': bool(field.get('required')),
            })

    bloc_context = get_public_bloc_context(event, include_inactive=False)
    bloc_columns = []
    if bloc_context:
        header_by_key = {
            'status': STATUS_HEADER,
            'restauration': RESTAURATION_HEADER,
            'social_event': SOCIAL_EVENT_HEADER,
        }
        for bloc in bloc_context['custom_blocs']:
            bloc_columns.append({
                'header': header_by_key[bloc['key']],
                'bloc_key': bloc['key'],
                'required': bloc['key'] == 'status',
                'select_mode': bloc['select_mode'],
                'valid_values': [item.name for item in bloc['items']],
            })
        if bloc_context['paid_sessions']:
            bloc_columns.append({
                'header': WORKSHOPS_HEADER,
                'bloc_key': 'workshops',
                'required': False,
                'select_mode': 'multiple',
                'valid_values': [s.title for s in bloc_context['paid_sessions']],
            })

    return {'form_fields': form_fields, 'bloc_columns': bloc_columns, 'bloc_context': bloc_context}


def generate_import_template_workbook(event):
    """Returns an openpyxl Workbook: an 'Import' sheet with headers (+ a sample
    row), and a 'Valeurs valides' sheet listing every bloc/workshop column's
    exact accepted values -- so whoever fills the sheet can copy-paste
    the right spelling instead of guessing."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    columns = build_import_columns(event)
    wb = Workbook()

    ws = wb.active
    ws.title = 'Import'
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill('solid', fgColor='2D1B6B')

    headers = [f['label'] + (' *' if f['required'] else '') for f in columns['form_fields']]
    headers += [c['header'] for c in columns['bloc_columns']]
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        ws.column_dimensions[cell.column_letter].width = max(18, len(header) + 2)

    # One example row so the expected shape (comma-separated multi-values,
    # exact names) is obvious without needing the legend sheet open.
    example = []
    for f in columns['form_fields']:
        if f['type'] == 'email':
            example.append('jean.dupont@example.com')
        elif f['name'] in ('first_name', 'prenom'):
            example.append('Jean')
        elif f['name'] in ('last_name', 'nom'):
            example.append('Dupont')
        else:
            example.append('')
    for c in columns['bloc_columns']:
        if not c['valid_values']:
            example.append('')
        elif c['select_mode'] == 'multiple':
            example.append(', '.join(c['valid_values'][:2]))
        else:
            example.append(c['valid_values'][0])
    ws.append(example)

    if columns['bloc_columns']:
        legend = wb.create_sheet('Valeurs valides')
        legend.append(['Colonne', 'Valeurs acceptées (orthographe exacte)'])
        legend['A1'].font = Font(bold=True)
        legend['B1'].font = Font(bold=True)
        for c in columns['bloc_columns']:
            note = (
                'Une seule valeur' if c['select_mode'] == 'single'
                else 'Une ou plusieurs valeurs, séparées par des virgules'
            )
            legend.append([c['header'], note])
            for value in c['valid_values']:
                legend.append(['', value])
        legend.column_dimensions['A'].width = 30
        legend.column_dimensions['B'].width = 50

    return wb


def _match_names(raw_value, valid_values, allow_multiple):
    """
    Splits a cell's raw text (comma-separated if allow_multiple) and
    case/whitespace-insensitively matches each piece against
    valid_values. Returns (matched_values: list, unmatched: list) --
    matched_values preserves each valid_values' own canonical spelling,
    never the user's raw typed text.
    """
    if not raw_value or not str(raw_value).strip():
        return [], []
    pieces = [p.strip() for p in str(raw_value).split(',')] if allow_multiple else [str(raw_value).strip()]
    pieces = [p for p in pieces if p]
    lookup = {v.lower(): v for v in valid_values}
    matched, unmatched = [], []
    for piece in pieces:
        canonical = lookup.get(piece.lower())
        if canonical:
            matched.append(canonical)
        else:
            unmatched.append(piece)
    return matched, unmatched


def parse_uploaded_workbook(uploaded_file):
    """Reads the first sheet's header row + data rows into a list of
    {header: raw_cell_value} dicts, 1:1 with build_import_columns'
    headers. Blank trailing rows are skipped."""
    from openpyxl import load_workbook

    wb = load_workbook(uploaded_file, data_only=True, read_only=True)
    ws = wb.worksheets[0]
    rows_iter = ws.iter_rows(values_only=True)
    header_row = next(rows_iter, None)
    if not header_row:
        return []

    headers = [str(h).strip() if h is not None else '' for h in header_row]
    parsed = []
    for raw_row in rows_iter:
        if raw_row is None or all(cell is None or str(cell).strip() == '' for cell in raw_row):
            continue
        row_dict = {}
        for idx, header in enumerate(headers):
            if not header:
                continue
            row_dict[header] = raw_row[idx] if idx < len(raw_row) else None
        parsed.append(row_dict)
    return parsed


def _validate_row(event, columns, row_dict):
    """
    Returns a dict ready to become a ParticipantImportRow's fields
    (email, full_name, form_data, resolved_item_ids, resolved_session_ids,
    status, error_message). Never raises -- an unexpected problem with
    one row becomes status='error' with a message, not a crash that
    takes the rest of the file down with it.
    """
    errors = []
    form_data = {}
    email = ''

    for field in columns['form_fields']:
        header = field['label'] + (' *' if field['required'] else '')
        raw_value = row_dict.get(header, '')
        value = '' if raw_value is None else str(raw_value).strip()
        if field['type'] == 'checkbox':
            value = [v.strip() for v in value.split(',') if v.strip()] if value else []
        if field['required'] and not value:
            errors.append(f"« {field['label']} » est requis.")
        form_data[field['name']] = value
        if field['type'] == 'email' or field['name'] in ('email',):
            email = (value or '').lower()

    if not email:
        errors.append("E-mail manquant ou invalide.")

    first_name = form_data.get('first_name') or form_data.get('prenom') or ''
    last_name = form_data.get('last_name') or form_data.get('nom') or ''
    full_name = f"{first_name} {last_name}".strip()

    resolved_item_ids = []
    resolved_session_ids = []
    for col in columns['bloc_columns']:
        raw_value = row_dict.get(col['header'], '')
        allow_multiple = col['select_mode'] == 'multiple'
        matched, unmatched = _match_names(raw_value, col['valid_values'], allow_multiple)

        if col['required'] and not matched:
            errors.append(f"« {col['header']} » est requis et doit correspondre exactement à une valeur de la feuille « Valeurs valides ».")
        if unmatched:
            errors.append(f"« {col['header']} » : valeur(s) non reconnue(s) : {', '.join(unmatched)}.")
        if not allow_multiple and len(matched) > 1:
            errors.append(f"« {col['header']} » n'accepte qu'une seule valeur.")

        if col['bloc_key'] == 'workshops':
            sessions_by_name = {s.title: s for s in columns['bloc_context']['paid_sessions']}
            resolved_session_ids.extend(str(sessions_by_name[name].id) for name in matched if name in sessions_by_name)
        else:
            items_by_name = {
                item.name: item for b in columns['bloc_context']['custom_blocs'] if b['key'] == col['bloc_key']
                for item in b['items']
            }
            resolved_item_ids.extend(str(items_by_name[name].id) for name in matched if name in items_by_name)

    if email and not errors:
        from events.models import ParticipantEventRegistration
        already_registered = ParticipantEventRegistration.objects.filter(
            event=event, participant__user__email__iexact=email
        ).exists()
        if already_registered:
            return {
                'email': email, 'full_name': full_name, 'form_data': form_data,
                'resolved_item_ids': resolved_item_ids, 'resolved_session_ids': resolved_session_ids,
                'status': 'skipped_duplicate', 'error_message': "Déjà inscrit à cet événement -- ignoré.",
            }

    return {
        'email': email, 'full_name': full_name, 'form_data': form_data,
        'resolved_item_ids': resolved_item_ids, 'resolved_session_ids': resolved_session_ids,
        'status': 'error' if errors else 'valid',
        'error_message': ' '.join(errors),
    }


def create_import_batch(event, user, uploaded_file):
    """Parses + validates the uploaded file and persists a
    ParticipantImportBatch with one ParticipantImportRow per data row,
    ready for the admin to review before anything is actually created."""
    columns = build_import_columns(event)
    parsed_rows = parse_uploaded_workbook(uploaded_file)

    batch = ParticipantImportBatch.objects.create(
        event=event, uploaded_by=user,
        original_filename=getattr(uploaded_file, 'name', '')[:255],
        status='ready',
    )

    rows = []
    for row_number, row_dict in enumerate(parsed_rows, start=2):  # row 1 is the header
        validated = _validate_row(event, columns, row_dict)
        rows.append(ParticipantImportRow(
            batch=batch, row_number=row_number, raw_data={k: ('' if v is None else str(v)) for k, v in row_dict.items()},
            **validated,
        ))
    ParticipantImportRow.objects.bulk_create(rows)

    return batch


def process_import_row(row, event, config, bloc_context):
    """
    Creates the real registration for one 'valid' row -- same pipeline
    as event_owner_new_registration: passwordless user, Participant +
    ParticipantEventRegistration, a FormSubmission capturing the raw
    form answers, and (only if this event has blocs at all) a
    RegistrationOrder priced through compute_order exactly like a real
    submission, left at its default 'pending' (unpaid/"Registered")
    status -- never auto-marked paid just because it came from a file.
    Mutates row.status/error_message/participant; caller saves it.
    """
    from .models_form import FormSubmission

    try:
        with transaction.atomic():
            first_name = row.form_data.get('first_name') or row.form_data.get('prenom') or ''
            last_name = row.form_data.get('last_name') or row.form_data.get('nom') or ''

            user = get_or_create_user_for_manual_registration(row.email, first_name, last_name)
            participant = create_participant_for_event(user, event)

            form_config = event.custom_forms.first()
            submission = None
            if form_config:
                submission = FormSubmission.objects.create(
                    form=form_config, data=row.form_data, email=row.email,
                )

            if bloc_context:
                result = compute_order(
                    event=event, config=config,
                    selected_item_ids=row.resolved_item_ids, selected_session_ids=row.resolved_session_ids,
                    on_date=timezone.now().date(),
                )
                from .models_blocs import RegistrationOrder
                RegistrationOrder.objects.create(
                    event=event, form_submission=submission, participant=participant,
                    period_id=result['active_period_id'], full_name=row.full_name, email=row.email,
                    items_snapshot=result['snapshot'], subtotals=result['subtotals'],
                    distinct_blocs_count=result['distinct_blocs_count'],
                    total_before_reduction=result['total_before_reduction'],
                    period_discount_percent=result['period_discount_percent'],
                    blocs_discount_percent=result['blocs_discount_percent'],
                    total_discount_percent=result['total_discount_percent'],
                    total_after_reduction=result['total_after_reduction'],
                    receipt_file='',  # bulk-imported rows never carry an upload receipt
                )

            row.status = 'created'
            row.participant = participant
            row.error_message = ''
    except Exception as e:
        row.status = 'error'
        row.error_message = f"Erreur lors de la création : {e}"
