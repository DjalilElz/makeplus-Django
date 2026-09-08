"""
Views for displaying email campaign and form analytics statistics.
Provides Brevo-like interface for viewing detailed stats.
"""
from django.shortcuts import render, get_object_or_404, redirect
from django.contrib.auth.decorators import login_required, user_passes_test
from django.db.models import Count, Q, Avg, Sum, F
from django.db.models.functions import Trunc
from django.http import HttpResponse
from django.utils import timezone
from datetime import timedelta, datetime, time
from .models_email import EmailCampaign, EmailRecipient, EmailLink, EmailClick, EmailOpen
from .models_form import FormConfiguration, FormAnalytics, FormView, FormFieldInteraction
from .views import is_staff_user


@login_required
def campaign_stats_detail(request, campaign_id):
    """Detailed statistics for an email campaign (Brevo-style)"""
    campaign = get_object_or_404(EmailCampaign, id=campaign_id)
    
    # Overall statistics
    total_recipients = campaign.recipients.count()
    # Count 'sent' and 'delivered' as delivered
    delivered = campaign.recipients.filter(status__in=['sent', 'delivered']).count()
    bounced = campaign.recipients.filter(status='bounced').count()
    unsubscribed = campaign.recipients.filter(status='unsubscribed').count()
    
    # Engagement stats - recalculate from actual opens
    unique_opens = campaign.recipients.filter(open_count__gt=0).count()
    unique_clicks = campaign.recipients.filter(click_count__gt=0).count()
    
    # Calculate rates based on delivered
    open_rate = round((unique_opens / delivered * 100) if delivered > 0 else 0, 2)
    click_rate = round((unique_clicks / delivered * 100) if delivered > 0 else 0, 2)
    
    # Click-to-open rate (CTOR)
    ctor = round((unique_clicks / unique_opens * 100) if unique_opens > 0 else 0, 2)
    
    # Get top performing links
    top_links = EmailLink.objects.filter(campaign=campaign).order_by('-unique_clicks')[:10]
    
    # Get most engaged recipients
    top_recipients = EmailRecipient.objects.filter(
        campaign=campaign
    ).annotate(
        engagement_score=F('open_count') + (F('click_count') * 2)
    ).order_by('-engagement_score')[:20]
    
    # Detailed recipient lists
    all_recipients_list = campaign.recipients.all().order_by('-open_count', '-click_count')
    
    # Recipients who opened (at least once)
    recipients_who_opened = campaign.recipients.filter(open_count__gt=0).order_by('-open_count')
    
    # Recipients who clicked (at least once) with unique links count
    recipients_who_clicked = []
    for recipient in campaign.recipients.filter(click_count__gt=0).order_by('-click_count'):
        # Count unique links this recipient clicked
        unique_links = EmailClick.objects.filter(recipient=recipient).values('link').distinct().count()
        recipient.unique_links_clicked = unique_links
        recipients_who_clicked.append(recipient)
    
    # Recipients who did not open
    recipients_not_opened = campaign.recipients.filter(open_count=0).order_by('email')
    
    # Timeline data - opens and clicks over time
    opens_timeline = EmailOpen.objects.filter(
        recipient__campaign=campaign
    ).annotate(
        hour=Trunc('opened_at', 'hour')
    ).values('hour').annotate(count=Count('id')).order_by('hour')
    
    clicks_timeline = EmailClick.objects.filter(
        recipient__campaign=campaign
    ).annotate(
        hour=Trunc('clicked_at', 'hour')
    ).values('hour').annotate(count=Count('id')).order_by('hour')
    
    # Device/client breakdown (from user agents)
    device_stats = {
        'desktop': 0,
        'mobile': 0,
        'tablet': 0,
        'unknown': 0
    }
    
    for recipient in campaign.recipients.all():
        opens = EmailOpen.objects.filter(recipient=recipient).first()
        if opens and opens.user_agent:
            ua_lower = opens.user_agent.lower()
            if 'mobile' in ua_lower or 'android' in ua_lower or 'iphone' in ua_lower:
                device_stats['mobile'] += 1
            elif 'tablet' in ua_lower or 'ipad' in ua_lower:
                device_stats['tablet'] += 1
            elif 'mozilla' in ua_lower or 'chrome' in ua_lower or 'safari' in ua_lower:
                device_stats['desktop'] += 1
            else:
                device_stats['unknown'] += 1
    
    # Geographic data (top locations by IP)
    locations = EmailOpen.objects.filter(
        recipient__campaign=campaign,
        ip_address__isnull=False
    ).values('ip_address').annotate(
        count=Count('id')
    ).order_by('-count')[:10]
    
    context = {
        'campaign': campaign,
        'total_recipients': total_recipients,
        'delivered': delivered,
        'bounced': bounced,
        'unsubscribed': unsubscribed,
        'unique_opens': unique_opens,
        'unique_clicks': unique_clicks,
        'open_rate': open_rate,
        'click_rate': click_rate,
        'ctor': ctor,
        'top_links': top_links,
        'top_recipients': top_recipients,
        'all_recipients_list': all_recipients_list,
        'recipients_who_opened': recipients_who_opened,
        'recipients_who_clicked': recipients_who_clicked,
        'recipients_not_opened': recipients_not_opened,
        'opens_timeline': list(opens_timeline),
        'clicks_timeline': list(clicks_timeline),
        'device_stats': device_stats,
        'locations': locations,
    }
    
    response = render(request, 'dashboard/campaign_stats_detail.html', context)
    # Add cache control headers to prevent caching
    response['Cache-Control'] = 'no-cache, no-store, must-revalidate, max-age=0'
    response['Pragma'] = 'no-cache'
    response['Expires'] = '0'
    
    return response


@login_required
def campaign_recipient_detail(request, campaign_id, recipient_id):
    """Detailed view of a single recipient's engagement"""
    campaign = get_object_or_404(EmailCampaign, id=campaign_id)
    recipient = get_object_or_404(EmailRecipient, id=recipient_id, campaign=campaign)
    
    # Get all opens
    opens = EmailOpen.objects.filter(recipient=recipient).order_by('-opened_at')
    
    # Get all clicks
    clicks = EmailClick.objects.filter(recipient=recipient).select_related('link').order_by('-clicked_at')
    
    # Get link summary - which links were clicked and how many times
    links_clicked = EmailClick.objects.filter(recipient=recipient).values('link').distinct()
    links_clicked_summary = []
    for link_data in links_clicked:
        link = EmailLink.objects.get(id=link_data['link'])
        click_events = EmailClick.objects.filter(recipient=recipient, link=link).order_by('clicked_at')
        links_clicked_summary.append({
            'url': link.original_url,
            'click_count': click_events.count(),
            'first_click': click_events.first().clicked_at if click_events.exists() else None,
            'last_click': click_events.last().clicked_at if click_events.exists() else None,
        })
    
    # Sort by click count descending
    links_clicked_summary = sorted(links_clicked_summary, key=lambda x: x['click_count'], reverse=True)
    
    context = {
        'campaign': campaign,
        'recipient': recipient,
        'opens': opens,
        'clicks': clicks,
        'links_clicked_summary': links_clicked_summary,
    }
    
    return render(request, 'dashboard/campaign_recipient_detail.html', context)


@login_required
def form_stats_detail(request, form_id):
    """Detailed statistics for a form (Brevo-style)"""
    form = get_object_or_404(FormConfiguration, id=form_id)
    
    # Get or create analytics
    analytics, created = FormAnalytics.objects.get_or_create(form=form)
    
    # Overall stats
    total_views = analytics.total_views
    total_submissions = analytics.total_submissions
    conversion_rate = analytics.conversion_rate
    
    # Device breakdown
    device_breakdown = analytics.device_breakdown or {}
    desktop_views = device_breakdown.get('desktop', 0)
    mobile_views = device_breakdown.get('mobile', 0)
    tablet_views = device_breakdown.get('tablet', 0)
    
    # Traffic sources
    traffic_sources = analytics.traffic_sources or {}
    top_sources = sorted(traffic_sources.items(), key=lambda x: x[1], reverse=True)[:10]
    
    # Field-level analytics
    field_stats = FormFieldInteraction.objects.filter(
        form_view__form=form
    ).values('field_name').annotate(
        total_interactions=Count('id'),
        avg_time_spent=Avg('time_spent'),
        avg_changes=Avg('changes_count'),
        completion_rate=Count('id', filter=Q(completed=True)) * 100.0 / Count('id')
    ).order_by('-total_interactions')
    
    # Conversion funnel
    views_with_interaction = FormView.objects.filter(form=form).exclude(
        fieldinteractions__isnull=True
    ).distinct().count()
    
    started_rate = round((views_with_interaction / total_views * 100) if total_views > 0 else 0, 2)
    
    # Timeline - views and submissions over time
    views_timeline = FormView.objects.filter(form=form).extra(
        select={'date': "date(viewed_at)"}
    ).values('date').annotate(count=Count('id')).order_by('date')
    
    # Top UTM campaigns
    utm_campaigns = FormView.objects.filter(
        form=form,
        utm_campaign__isnull=False
    ).values('utm_campaign').annotate(
        views=Count('id'),
        conversions=Count('id', filter=Q(completed=True))
    ).order_by('-views')[:10]
    
    # Browser stats
    browser_stats = FormView.objects.filter(
        form=form,
        browser__isnull=False
    ).values('browser').annotate(
        count=Count('id')
    ).order_by('-count')[:10]
    
    # Average time on form
    avg_time_stats = FormFieldInteraction.objects.filter(
        form_view__form=form
    ).aggregate(
        avg_time=Avg('time_spent')
    )
    
    # Dropout fields (fields with low completion rate)
    dropout_fields = FormFieldInteraction.objects.filter(
        form_view__form=form
    ).values('field_name').annotate(
        started=Count('id'),
        completed=Count('id', filter=Q(completed=True)),
        dropout_rate=(Count('id') - Count('id', filter=Q(completed=True))) * 100.0 / Count('id')
    ).order_by('-dropout_rate')[:10]
    
    context = {
        'form': form,
        'analytics': analytics,
        'total_views': total_views,
        'total_submissions': total_submissions,
        'conversion_rate': conversion_rate,
        'started_rate': started_rate,
        'desktop_views': desktop_views,
        'mobile_views': mobile_views,
        'tablet_views': tablet_views,
        'top_sources': top_sources,
        'field_stats': field_stats,
        'views_timeline': list(views_timeline),
        'utm_campaigns': utm_campaigns,
        'browser_stats': browser_stats,
        'avg_time': avg_time_stats['avg_time'] or 0,
        'dropout_fields': dropout_fields,
    }
    
    return render(request, 'dashboard/form_stats_detail.html', context)


@login_required
def campaign_list_with_stats(request):
    """List all campaigns with summary stats"""
    from events.models import Event
    from django.views.decorators.cache import never_cache
    
    # Get filter parameters
    status_filter = request.GET.get('status', '')
    event_filter = request.GET.get('event', '')
    
    # Base queryset
    campaigns = EmailCampaign.objects.select_related('event').order_by('-created_at')
    
    # Apply filters
    if status_filter:
        campaigns = campaigns.filter(status=status_filter)
    if event_filter:
        campaigns = campaigns.filter(event_id=event_filter)
    
    # Get all events for filter dropdown
    events = Event.objects.all().order_by('-created_at')
    
    # Calculate counts for filter buttons
    all_campaigns = EmailCampaign.objects.all()
    total_count = all_campaigns.count()
    draft_count = all_campaigns.filter(status='draft').count()
    sent_count = all_campaigns.filter(status='sent').count()
    sending_count = all_campaigns.filter(status='sending').count()
    
    # Add stats to each campaign
    campaign_stats = []
    for campaign in campaigns:
        recipients = campaign.recipients.all()
        total_recipients = recipients.count()
        sent_count_campaign = recipients.filter(status__in=['sent', 'delivered']).count()
        opened_count = recipients.filter(open_count__gt=0).count()
        clicked_count = recipients.filter(click_count__gt=0).count()
        
        campaign_stats.append({
            'campaign': campaign,
            'total_recipients': total_recipients,
            'sent_count': sent_count_campaign,
            'opened_count': opened_count,
            'clicked_count': clicked_count,
            'open_rate': campaign.get_open_rate(),
            'click_rate': campaign.get_click_rate(),
            'status': campaign.status,
        })
    
    context = {
        'campaign_stats': campaign_stats,
        'events': events,
        'total_count': total_count,
        'draft_count': draft_count,
        'sent_count': sent_count,
        'sending_count': sending_count,
    }
    
    response = render(request, 'dashboard/campaign_list_with_stats.html', context)
    # Add cache control headers to prevent caching
    response['Cache-Control'] = 'no-cache, no-store, must-revalidate, max-age=0'
    response['Pragma'] = 'no-cache'
    response['Expires'] = '0'
    
    return response


@login_required
def form_list_with_stats(request):
    """List all forms with summary stats"""
    forms = FormConfiguration.objects.select_related('event', 'created_by').order_by('-created_at')
    
    # Add analytics to each form
    form_stats = []
    for form in forms:
        try:
            analytics = FormAnalytics.objects.get(form=form)
        except FormAnalytics.DoesNotExist:
            analytics = FormAnalytics.objects.create(form=form)
        
        form_stats.append({
            'form': form,
            'analytics': analytics,
        })
    
    context = {
        'form_stats': form_stats,
    }

    return render(request, 'dashboard/form_list_with_stats.html', context)


# ==================== Event Stats (combined caisse/attendance/scans) ====================

# Dashboard tables here (transactions, presence, scans) are display-only
# overviews, not exports -- capped so a large event's history can't make
# this page itself slow to render.
_EVENT_STATS_LIST_LIMIT = 200


def _parse_stats_date_range(request):
    """
    (date_from, date_to, date_from_str, date_to_str) from ?date_from=&date_to=,
    submitted by an <input type="datetime-local"> so the filter can narrow
    to an exact hour/minute ("between X and Y"), not just a whole day.
    date_from/date_to come back as exact timezone-aware instants (no
    start/end-of-day padding) since the input already carries a time.
    A bare YYYY-MM-DD (an older bookmarked link, or a browser that only
    submits a date) is still accepted and expanded to the full day so
    those links keep working. Invalid/missing values fall back to no
    bound (the empty string is also what re-populates the form's own
    inputs).
    """
    date_from_str = request.GET.get('date_from', '').strip()
    date_to_str = request.GET.get('date_to', '').strip()
    date_from = None
    date_to = None
    if date_from_str:
        try:
            date_from = timezone.make_aware(datetime.strptime(date_from_str, '%Y-%m-%dT%H:%M'))
        except ValueError:
            try:
                date_from = timezone.make_aware(
                    datetime.combine(datetime.strptime(date_from_str, '%Y-%m-%d').date(), time.min)
                )
            except ValueError:
                date_from_str = ''
    if date_to_str:
        try:
            date_to = timezone.make_aware(datetime.strptime(date_to_str, '%Y-%m-%dT%H:%M'))
        except ValueError:
            try:
                date_to = timezone.make_aware(
                    datetime.combine(datetime.strptime(date_to_str, '%Y-%m-%d').date(), time.max)
                )
            except ValueError:
                date_to_str = ''
    return date_from, date_to, date_from_str, date_to_str


def _in_date_range(queryset, field, date_from, date_to):
    if date_from:
        queryset = queryset.filter(**{f'{field}__gte': date_from})
    if date_to:
        queryset = queryset.filter(**{f'{field}__lte': date_to})
    return queryset


def _gather_event_stats(request, event):
    """
    Shared filter + query logic for the Stats page and its PDF/Excel
    exports, so an export always matches exactly what's on screen for the
    same filters (date range/hour, caisse, presence search, scan search).
    Returns unsliced querysets -- the page caps list sizes for render
    speed (_EVENT_STATS_LIST_LIMIT), the exports don't, since an export is
    explicitly asked for the full detail.
    """
    from events.models import ParticipantEventRegistration, ControllerScan, ExposantScan
    from caisse.models import Caisse, CaisseTransaction

    date_from, date_to, date_from_str, date_to_str = _parse_stats_date_range(request)
    scan_search = request.GET.get('scan_search', '').strip()
    presence_search = request.GET.get('presence_search', '').strip()

    caisses = list(Caisse.objects.filter(event=event).order_by('name'))
    caisse_id = request.GET.get('caisse_id', '').strip()
    selected_caisse = None
    if caisse_id:
        selected_caisse = next((c for c in caisses if str(c.id) == caisse_id), None)
        if not selected_caisse:
            caisse_id = ''  # unknown/stale id -- fall back to "all caisses"

    # ---- Money & transactions (every caisse station combined, unless one
    # is selected in the filter bar) ----
    # NOTE: the per-caisse .values().annotate() grouping below must run
    # against an UN-ordered queryset -- Django folds order_by() fields into
    # GROUP BY for a .values().annotate() call, which would silently break
    # the grouping (one row per transaction instead of per caisse) if
    # '-created_at' were applied first. The list(...) call materializes it
    # immediately, so reordering the queryset afterwards for display is safe.
    transactions_base_qs = CaisseTransaction.objects.filter(caisse__event=event, status='completed')
    if selected_caisse:
        transactions_base_qs = transactions_base_qs.filter(caisse=selected_caisse)
    transactions_base_qs = _in_date_range(transactions_base_qs, 'created_at', date_from, date_to)

    money_stats = transactions_base_qs.aggregate(
        total_amount=Sum('total_amount'),
        transaction_count=Count('id'),
        total_participants=Count('participant_id', distinct=True),
    )
    per_caisse_stats = list(transactions_base_qs.values('caisse_id', 'caisse__name').annotate(
        total_amount=Sum('total_amount'),
        transaction_count=Count('id'),
        total_participants=Count('participant_id', distinct=True),
    ).order_by('caisse__name'))

    transactions_qs = transactions_base_qs.select_related(
        'caisse', 'participant__user',
    ).prefetch_related('items').order_by('-created_at')

    # ---- Presence (event-wide check-in, from any caisse or badge scan),
    # optionally narrowed to a participant by name/e-mail/badge ----
    presence_qs = _in_date_range(
        ParticipantEventRegistration.objects.filter(event=event, is_checked_in=True),
        'checked_in_at', date_from, date_to,
    ).select_related('participant__user')
    if presence_search:
        presence_qs = presence_qs.filter(
            Q(participant__user__first_name__icontains=presence_search)
            | Q(participant__user__last_name__icontains=presence_search)
            | Q(participant__user__email__icontains=presence_search)
            | Q(participant__badge_id__icontains=presence_search)
        )
    presence_qs = presence_qs.order_by('-checked_in_at')
    presence_count = presence_qs.count()
    total_registered = ParticipantEventRegistration.objects.filter(event=event).count()

    # ---- Scans (badge controllers + exhibitor booth visits), filterable
    # by the same date range, by a specific scanned participant, and by
    # WHICH controller/exposant did the scanning ----
    controller_id = request.GET.get('controller_id', '').strip()
    exposant_id = request.GET.get('exposant_id', '').strip()

    # Dropdown options are the controllers/exposants who actually have scan
    # history for this event (not every assigned controller/exposant) --
    # anyone with zero scans has nothing to filter down to anyway.
    controller_options = list(
        ControllerScan.objects.filter(event=event).select_related('controller')
        .values('controller_id', 'controller__first_name', 'controller__last_name', 'controller__username')
        .distinct().order_by('controller__first_name', 'controller__last_name')
    )
    exposant_options = list(
        ExposantScan.objects.filter(event=event).select_related('exposant__user')
        .values('exposant_id', 'exposant__user__first_name', 'exposant__user__last_name', 'exposant__user__username')
        .distinct().order_by('exposant__user__first_name', 'exposant__user__last_name')
    )

    controller_scans_qs = _in_date_range(
        ControllerScan.objects.filter(event=event), 'scanned_at', date_from, date_to,
    ).select_related('controller')
    exposant_scans_qs = _in_date_range(
        ExposantScan.objects.filter(event=event), 'scanned_at', date_from, date_to,
    ).select_related('exposant__user', 'scanned_participant__user')

    if controller_id:
        controller_scans_qs = controller_scans_qs.filter(controller_id=controller_id)
    if exposant_id:
        exposant_scans_qs = exposant_scans_qs.filter(exposant_id=exposant_id)

    if scan_search:
        controller_scans_qs = controller_scans_qs.filter(
            Q(participant_name__icontains=scan_search)
            | Q(participant_email__icontains=scan_search)
            | Q(badge_id__icontains=scan_search)
        )
        exposant_scans_qs = exposant_scans_qs.filter(
            Q(scanned_participant__user__first_name__icontains=scan_search)
            | Q(scanned_participant__user__last_name__icontains=scan_search)
            | Q(scanned_participant__user__email__icontains=scan_search)
            | Q(scanned_participant__badge_id__icontains=scan_search)
        )

    controller_scans_qs = controller_scans_qs.order_by('-scanned_at')
    exposant_scans_qs = exposant_scans_qs.order_by('-scanned_at')

    return {
        'date_from_str': date_from_str, 'date_to_str': date_to_str,
        'scan_search': scan_search, 'presence_search': presence_search,
        'caisses': caisses, 'caisse_id': caisse_id, 'selected_caisse': selected_caisse,
        'money_stats': money_stats, 'per_caisse_stats': per_caisse_stats,
        'transactions_qs': transactions_qs,
        'presence_qs': presence_qs, 'presence_count': presence_count, 'total_registered': total_registered,
        'controller_scans_qs': controller_scans_qs, 'exposant_scans_qs': exposant_scans_qs,
        'controller_id': controller_id, 'exposant_id': exposant_id,
        'controller_options': controller_options, 'exposant_options': exposant_options,
    }


@login_required
@user_passes_test(is_staff_user)
def event_stats(request, event_id):
    """
    Combined, date-filterable view of everything happening at an event on
    the ground: money collected across ALL caisses (not just one
    station's own dashboard), who's actually present, and scan history --
    previously scattered with no single place to see it all together or
    narrow it to a specific day/hour range.
    """
    from events.models import Event

    event = get_object_or_404(Event, id=event_id)
    data = _gather_event_stats(request, event)

    transactions_qs = data['transactions_qs']
    transaction_count = data['money_stats']['transaction_count'] or 0
    controller_scans_qs = data['controller_scans_qs']
    exposant_scans_qs = data['exposant_scans_qs']

    context = {
        'event': event,
        'date_from': data['date_from_str'],
        'date_to': data['date_to_str'],
        'scan_search': data['scan_search'],
        'presence_search': data['presence_search'],
        'controller_id': data['controller_id'],
        'exposant_id': data['exposant_id'],
        'controller_options': data['controller_options'],
        'exposant_options': data['exposant_options'],
        'caisses': data['caisses'],
        'caisse_id': data['caisse_id'],
        'selected_caisse': data['selected_caisse'],
        'money_stats': data['money_stats'],
        'per_caisse_stats': data['per_caisse_stats'],
        'transactions': transactions_qs[:_EVENT_STATS_LIST_LIMIT],
        'transactions_total_count': transaction_count,
        'transactions_truncated': transaction_count > _EVENT_STATS_LIST_LIMIT,
        'presence_list': data['presence_qs'][:_EVENT_STATS_LIST_LIMIT],
        'presence_count': data['presence_count'],
        'presence_truncated': data['presence_count'] > _EVENT_STATS_LIST_LIMIT,
        'total_registered': data['total_registered'],
        'controller_scans': controller_scans_qs[:_EVENT_STATS_LIST_LIMIT],
        'controller_scans_count': controller_scans_qs.count(),
        'exposant_scans': exposant_scans_qs[:_EVENT_STATS_LIST_LIMIT],
        'exposant_scans_count': exposant_scans_qs.count(),
        'export_querystring': request.GET.urlencode(),
    }
    return render(request, 'dashboard/event_stats.html', context)


def _stats_period_label(data):
    return f"{data['date_from_str'] or 'depuis le début'} → {data['date_to_str'] or 'aujourd’hui'}"


@login_required
@user_passes_test(is_staff_user)
def event_stats_export_excel(request, event_id):
    """
    Full-detail Excel export of the Stats page for the exact filters
    currently applied -- one workbook, one sheet per tab, uncapped (the
    on-page tables cap at _EVENT_STATS_LIST_LIMIT for render speed; an
    export is explicitly asked for everything).
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from events.models import Event

    event = get_object_or_404(Event, id=event_id)
    data = _gather_event_stats(request, event)
    money_stats = data['money_stats']

    header_fill = PatternFill(start_color="2D1B6B", end_color="2D1B6B", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=12)
    border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin'),
    )

    def _write_header(ws, row_num, headers):
        for col, label in enumerate(headers, start=1):
            cell = ws.cell(row=row_num, column=col, value=label)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = border

    def _write_sheet(ws, headers, rows):
        _write_header(ws, 1, headers)
        for row in rows:
            ws.append(row)
        for col in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col)].width = 24

    wb = Workbook()

    # ---- Aperçu ----
    ws = wb.active
    ws.title = "Aperçu"
    ws.append(["Événement", event.name])
    ws.append(["Période", _stats_period_label(data)])
    ws.append(["Caisse", data['selected_caisse'].name if data['selected_caisse'] else 'Toutes les caisses'])
    ws.append([])
    ws.append(["Montant total collecté (DZD)", float(money_stats['total_amount'] or 0)])
    ws.append(["Transactions", money_stats['transaction_count'] or 0])
    ws.append(["Participants traités", money_stats['total_participants'] or 0])
    ws.append(["Participants présents", data['presence_count']])
    ws.append(["Total inscrits", data['total_registered']])
    ws.append([])
    header_row = ws.max_row + 1
    _write_header(ws, header_row, ["Caisse", "Montant (DZD)", "Transactions", "Participants"])
    total_amount = 0.0
    total_txn = 0
    for row in data['per_caisse_stats']:
        amount = float(row['total_amount'] or 0)
        ws.append([row['caisse__name'], amount, row['transaction_count'], row['total_participants']])
        total_amount += amount
        total_txn += row['transaction_count']
    ws.append(["TOTAL", total_amount, total_txn, money_stats['total_participants'] or 0])
    for col in range(1, 5):
        ws.cell(row=ws.max_row, column=col).font = Font(bold=True)
        ws.column_dimensions[get_column_letter(col)].width = 26

    # ---- Transactions ----
    rows = []
    for txn in data['transactions_qs']:
        items = ', '.join(i.name for i in txn.items.all())
        rows.append([
            txn.participant.user.get_full_name() or txn.participant.user.username,
            txn.participant.user.email,
            txn.caisse.name,
            items,
            float(txn.total_amount),
            txn.get_payment_method_display(),
            timezone.localtime(txn.created_at).strftime('%d/%m/%Y %H:%M'),
        ])
    _write_sheet(
        wb.create_sheet("Transactions"),
        ["Participant", "E-mail", "Caisse", "Articles", "Montant (DZD)", "Méthode", "Heure"], rows,
    )

    # ---- Présence ----
    rows = []
    for reg in data['presence_qs']:
        rows.append([
            reg.participant.user.get_full_name() or reg.participant.user.username,
            reg.participant.user.email,
            reg.participant.badge_id,
            timezone.localtime(reg.checked_in_at).strftime('%d/%m/%Y %H:%M') if reg.checked_in_at else '',
        ])
    _write_sheet(wb.create_sheet("Présence"), ["Participant", "E-mail", "Badge", "Présent depuis"], rows)

    # ---- Scans contrôleurs ----
    rows = []
    for scan in data['controller_scans_qs']:
        rows.append([
            scan.participant_name, scan.participant_email, scan.badge_id,
            scan.controller.get_full_name() or scan.controller.username,
            scan.get_status_display(),
            timezone.localtime(scan.scanned_at).strftime('%d/%m/%Y %H:%M'),
        ])
    _write_sheet(
        wb.create_sheet("Scans contrôleurs"),
        ["Participant", "E-mail", "Badge", "Contrôleur", "Statut", "Heure"], rows,
    )

    # ---- Scans exposants ----
    rows = []
    for scan in data['exposant_scans_qs']:
        rows.append([
            scan.scanned_participant.user.get_full_name() or scan.scanned_participant.user.username,
            scan.scanned_participant.user.email,
            scan.exposant.user.get_full_name() or scan.exposant.user.username,
            scan.notes or '',
            timezone.localtime(scan.scanned_at).strftime('%d/%m/%Y %H:%M'),
        ])
    _write_sheet(
        wb.create_sheet("Scans exposants"),
        ["Participant scanné", "E-mail", "Exposant", "Notes", "Heure"], rows,
    )

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="stats_{event.name}_{timezone.now().date()}.xlsx"'
    wb.save(response)
    return response


@login_required
@user_passes_test(is_staff_user)
def event_stats_export_pdf(request, event_id):
    """
    Full-detail PDF export of the Stats page, same filters/data as the
    Excel export -- landscape A4 so the wider tables (transactions, scans)
    stay readable; each table auto-paginates across pages.
    """
    from io import BytesIO
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import mm
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
    from events.models import Event

    event = get_object_or_404(Event, id=event_id)
    data = _gather_event_stats(request, event)
    money_stats = data['money_stats']
    styles = getSampleStyleSheet()

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(A4),
        topMargin=15 * mm, bottomMargin=15 * mm, leftMargin=12 * mm, rightMargin=12 * mm,
    )
    elements = [
        Paragraph(f"Statistiques — {event.name}", styles['Title']),
        Paragraph(
            f"Période : {_stats_period_label(data)} &nbsp;&nbsp;|&nbsp;&nbsp; "
            f"Caisse : {data['selected_caisse'].name if data['selected_caisse'] else 'Toutes les caisses'}",
            styles['Normal'],
        ),
        Spacer(1, 8),
    ]

    summary_table = Table([
        ["Montant total collecté", f"{float(money_stats['total_amount'] or 0):.2f} DZD"],
        ["Transactions", str(money_stats['transaction_count'] or 0)],
        ["Participants traités", str(money_stats['total_participants'] or 0)],
        ["Participants présents", f"{data['presence_count']} / {data['total_registered']}"],
    ], colWidths=[220, 160])
    summary_table.setStyle(TableStyle([
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#2D1B6B')),
        ('TEXTCOLOR', (0, 0), (0, -1), colors.white),
        ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
    ]))
    elements.append(summary_table)
    elements.append(Spacer(1, 14))

    def _section(title, headers, rows, empty_message):
        elements.append(Paragraph(title, styles['Heading2']))
        if not rows:
            elements.append(Paragraph(empty_message, styles['Italic']))
            elements.append(Spacer(1, 10))
            return
        table = Table([headers] + rows, repeatRows=1)
        table.setStyle(TableStyle([
            ('GRID', (0, 0), (-1, -1), 0.4, colors.grey),
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2D1B6B')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        elements.append(table)
        elements.append(Spacer(1, 14))

    per_caisse_rows = []
    total_amount = 0.0
    total_txn = 0
    for row in data['per_caisse_stats']:
        amount = float(row['total_amount'] or 0)
        per_caisse_rows.append([row['caisse__name'], f"{amount:.2f} DZD", str(row['transaction_count']), str(row['total_participants'])])
        total_amount += amount
        total_txn += row['transaction_count']
    per_caisse_rows.append(["TOTAL", f"{total_amount:.2f} DZD", str(total_txn), str(money_stats['total_participants'] or 0)])
    _section(
        "Détail par caisse", ["Caisse", "Montant", "Transactions", "Participants"],
        per_caisse_rows, "Aucune transaction pour cette période.",
    )

    elements.append(PageBreak())
    txn_rows = []
    for txn in data['transactions_qs']:
        items = ', '.join(i.name for i in txn.items.all())
        txn_rows.append([
            txn.participant.user.get_full_name() or txn.participant.user.username,
            txn.caisse.name, items, f"{float(txn.total_amount):.2f} DZD",
            txn.get_payment_method_display(),
            timezone.localtime(txn.created_at).strftime('%d/%m/%Y %H:%M'),
        ])
    _section(
        "Transactions", ["Participant", "Caisse", "Articles", "Montant", "Méthode", "Heure"],
        txn_rows, "Aucune transaction pour cette période.",
    )

    presence_rows = []
    for reg in data['presence_qs']:
        presence_rows.append([
            reg.participant.user.get_full_name() or reg.participant.user.username,
            reg.participant.badge_id,
            timezone.localtime(reg.checked_in_at).strftime('%d/%m/%Y %H:%M') if reg.checked_in_at else '',
        ])
    _section(
        "Présence", ["Participant", "Badge", "Présent depuis"],
        presence_rows, "Aucun participant présent pour cette période.",
    )

    controller_rows = []
    for scan in data['controller_scans_qs']:
        controller_rows.append([
            scan.participant_name, scan.badge_id,
            scan.controller.get_full_name() or scan.controller.username,
            scan.get_status_display(),
            timezone.localtime(scan.scanned_at).strftime('%d/%m/%Y %H:%M'),
        ])
    _section(
        "Scans des contrôleurs de badges", ["Participant", "Badge", "Contrôleur", "Statut", "Heure"],
        controller_rows, "Aucun scan pour ces critères.",
    )

    exposant_rows = []
    for scan in data['exposant_scans_qs']:
        exposant_rows.append([
            scan.scanned_participant.user.get_full_name() or scan.scanned_participant.user.username,
            scan.exposant.user.get_full_name() or scan.exposant.user.username,
            scan.notes or '',
            timezone.localtime(scan.scanned_at).strftime('%d/%m/%Y %H:%M'),
        ])
    _section(
        "Scans des exposants", ["Participant scanné", "Exposant", "Notes", "Heure"],
        exposant_rows, "Aucun scan pour ces critères.",
    )

    doc.build(elements)
    buffer.seek(0)
    response = HttpResponse(buffer.read(), content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="stats_{event.name}_{timezone.now().date()}.pdf"'
    return response
